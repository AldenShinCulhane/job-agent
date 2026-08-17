"""Resume upload -> raw text -> structured profile.

Text extraction handles PDF (pypdf) and DOCX (python-docx). A scanned or
image-based PDF yields near-empty text; that triggers an automatic OCR
fallback through ocrmypdf before giving up. Structured extraction is a
single medium-effort LLM call against PROFILE_SCHEMA.
"""

import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from backend.core import llm
from backend.resume.schema import EXTRACTION_SYSTEM_PROMPT, PROFILE_SCHEMA, extraction_prompt

# Below this many characters the extraction is considered failed/scanned.
_WEAK_TEXT_CHARS = 200
_OCR_TIMEOUT_S = 300


class IngestError(Exception):
    """User-fixable ingestion problem (bad file, OCR unavailable...)."""


def _pdf_text(path: Path) -> str:
    return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)


def _docx_text(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _is_weak(text: str) -> bool:
    return len(text.strip()) < _WEAK_TEXT_CHARS


def _ocr_pdf(path: Path) -> str:
    if shutil.which("ocrmypdf") is None:
        raise IngestError(
            "This PDF appears to be scanned (no extractable text), and "
            "ocrmypdf is not installed. Install it (macOS: brew install "
            "ocrmypdf | Ubuntu/Debian: sudo apt install ocrmypdf) or upload "
            "a text-based PDF/DOCX."
        )
    with tempfile.TemporaryDirectory() as tmp:
        ocred = Path(tmp) / "ocr.pdf"
        proc = subprocess.run(
            ["ocrmypdf", "--skip-text", str(path), str(ocred)],
            capture_output=True, text=True, timeout=_OCR_TIMEOUT_S,
        )
        if proc.returncode != 0 or not ocred.is_file():
            raise IngestError(f"OCR failed: {(proc.stderr or proc.stdout)[-500:]}")
        return _pdf_text(ocred)


def extract_text(path: Path) -> tuple[str, str]:
    """(text, method) where method is 'pdf' | 'pdf+ocr' | 'docx'."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        text = _pdf_text(path)
        if _is_weak(text):
            text = _ocr_pdf(path)
            if _is_weak(text):
                raise IngestError("Could not read text from this PDF, even with OCR.")
            return text, "pdf+ocr"
        return text, "pdf"
    if suffix == ".docx":
        text = _docx_text(path)
        if _is_weak(text):
            raise IngestError("This DOCX contains almost no text.")
        return text, "docx"
    raise IngestError(f"Unsupported file type {suffix!r} — upload a PDF or DOCX.")


async def extract_profile(resume_text: str) -> dict:
    """One structured extraction call (medium effort)."""
    return await llm.structured(
        extraction_prompt(resume_text),
        PROFILE_SCHEMA,
        task="extract",
        system_prompt=EXTRACTION_SYSTEM_PROMPT,
    )
